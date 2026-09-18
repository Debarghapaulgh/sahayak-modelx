"""
Create Master Excel Inspection Workbook for Structured-JSON SFT Part 1.
Builds SahayakAI_Part1_Structured_SFT_1500_Master.xlsx.
"""

import os
import json
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_structured_excel():
    base_dir = "datasets/textbook_sft_part1"
    
    files = {
        "QA_800": os.path.join(base_dir, "qa", "qa_structured_final.jsonl"),
        "Lesson_Plans_350": os.path.join(base_dir, "lesson_plans", "lesson_plan_structured_final.jsonl"),
        "Quizzes_350": os.path.join(base_dir, "quizzes", "quiz_structured_final.jsonl")
    }

    excel_path = "SahayakAI_Part1_Structured_SFT_1500_Master.xlsx"
    excel_path_mirror = os.path.join(base_dir, "SahayakAI_Part1_Structured_SFT_1500_Master.xlsx")

    dfs = {}
    for sheet_name, fpath in files.items():
        rows = []
        if os.path.exists(fpath):
            with open(fpath, "r", encoding="utf-8") as f:
                for line in f:
                    if not line.strip():
                        continue
                    r = json.loads(line)
                    user_prompt = ""
                    assistant_payload = {}
                    for m in r.get("messages", []):
                        if m.get("role") == "user":
                            user_prompt = m.get("content")
                        elif m.get("role") == "assistant":
                            assistant_payload = m.get("content")

                    rows.append({
                        "ID": r.get("id"),
                        "Task Type": r.get("task_type"),
                        "Requester": r.get("requester_role"),
                        "Target": r.get("instructional_target"),
                        "Board": r.get("curriculum", {}).get("board"),
                        "Stage": r.get("curriculum", {}).get("stage"),
                        "Grade": r.get("curriculum", {}).get("class"),
                        "Subject": r.get("curriculum", {}).get("subject"),
                        "Textbook": r.get("curriculum", {}).get("textbook"),
                        "Topic": r.get("curriculum", {}).get("topic"),
                        "User Prompt": user_prompt,
                        "Structured JSON Response": json.dumps(assistant_payload, ensure_ascii=False, indent=2),
                        "Validation Status": "APPROVED"
                    })
        dfs[sheet_name] = pd.DataFrame(rows)
        print(f"Loaded {len(rows)} rows for sheet {sheet_name}")

    for out_f in [excel_path, excel_path_mirror]:
        with pd.ExcelWriter(out_f, engine="openpyxl") as writer:
            for sheet_name, df in dfs.items():
                df.to_excel(writer, sheet_name=sheet_name, index=False)

        wb = openpyxl.load_workbook(out_f)
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10)
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            ws.views.sheetView[0].showGridLines = True

            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=1, column=col)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                ws.row_dimensions[1].height = 28

            for col in ws.columns:
                header_val = str(col[0].value)
                col_letter = get_column_letter(col[0].column)
                if "Response" in header_val:
                    ws.column_dimensions[col_letter].width = 65
                elif "Prompt" in header_val:
                    ws.column_dimensions[col_letter].width = 45
                elif "Topic" in header_val or "Textbook" in header_val:
                    ws.column_dimensions[col_letter].width = 25
                else:
                    ws.column_dimensions[col_letter].width = 14

        wb.save(out_f)
        print(f"Saved styled structured master Excel: {out_f}")

if __name__ == "__main__":
    build_structured_excel()